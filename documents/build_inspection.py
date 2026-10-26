"""Izvlači referencu za kazne i kontrolne liste iz Excel fajlova."""

from __future__ import annotations

from collections import OrderedDict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / 'Report Sloy B.xlsx'
LISTS = ROOT / 'Kontrolne liste digital empty.xlsx'

CYR = str.maketrans({
    'А': 'A', 'Б': 'B', 'В': 'V', 'Г': 'G', 'Д': 'D', 'Ђ': 'Đ', 'Е': 'E', 'Ж': 'Ž',
    'З': 'Z', 'И': 'I', 'Ј': 'J', 'К': 'K', 'Л': 'L', 'Љ': 'Lj', 'М': 'M', 'Н': 'N',
    'Њ': 'Nj', 'О': 'O', 'П': 'P', 'Р': 'R', 'С': 'S', 'Т': 'T', 'Ћ': 'Ć', 'У': 'U',
    'Ф': 'F', 'Х': 'H', 'Ц': 'C', 'Ч': 'Č', 'Џ': 'Dž', 'Ш': 'Š',
    'а': 'a', 'б': 'b', 'в': 'v', 'г': 'g', 'д': 'd', 'ђ': 'đ', 'е': 'e', 'ж': 'ž',
    'з': 'z', 'и': 'i', 'ј': 'j', 'к': 'k', 'л': 'l', 'љ': 'lj', 'м': 'm', 'н': 'n',
    'њ': 'nj', 'о': 'o', 'п': 'p', 'р': 'r', 'с': 's', 'т': 't', 'ћ': 'ć', 'у': 'u',
    'ф': 'f', 'х': 'h', 'ц': 'c', 'ч': 'č', 'џ': 'dž', 'ш': 'š',
})


def lat(text) -> str:
    if text is None:
        return ''
    return str(text).translate(CYR).strip()


def fines_body() -> str:
    wb = openpyxl.load_workbook(REPORT, data_only=True)
    rows: OrderedDict[str, tuple[str, str]] = OrderedDict()
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for row in ws.iter_rows(min_row=11, values_only=True):
            # kolone: nalaz | (prazno) | norma | kazna
            cells = list(row or ())
            finding = lat(cells[0] if cells else '')
            norm = lat(cells[2] if len(cells) > 2 else '')
            fine = lat(cells[3] if len(cells) > 3 else '')
            if not finding or finding.upper().startswith('UTVRĐENA'):
                continue
            key = finding.casefold()
            if key not in rows:
                rest = ' · '.join(p for p in (norm, fine) if p)
                rows[key] = (finding, rest)

    lines = [
        'Tipične neusaglašenosti koje se javljaju na eksternim proverama '
        'prema Zakonu o bezbednosti hrane (ZBH), sa okvirnim kaznenim merama.',
        '',
        'Izvor: Report Sloy B.xlsx (Food Compass) — sadržaj je uopšten kao '
        'referenca za pripremu; nije zapis o tvom objektu.',
        '',
        'Neusaglašenost · Pravna norma · Kaznena mera',
        '—',
    ]
    for i, (_, (finding, rest)) in enumerate(rows.items(), 1):
        lines.append(f'{i}. {finding}')
        lines.append(f'   {rest}')
        lines.append('')
    return '\n'.join(lines).strip() + '\n'


def checklist_body(sheet_name: str) -> str:
    wb = openpyxl.load_workbook(LISTS, data_only=True)
    ws = wb[sheet_name]
    title = lat(ws.cell(1, 1).value) or lat(sheet_name) or sheet_name
    lines = [
        f'Kontrolna lista: {title}',
        '',
        'Izvor: Kontrolne liste digital empty.xlsx (Food Compass).',
        'Odgovori tipično: DA / NE / NP. Koristi za pripremu pred inspekciju.',
        '',
    ]
    n = 0
    skip_prefixes = (
        'NAZIV SUBJEKTA', 'PITANJE', 'CHEKIRATI', 'ČEKIRATI', 'DA', 'NE',
        'NP', 'NAPOMENA', 'PRIMERBE', 'PRIMERBE', 'PRIMEDbe',
    )
    for row in ws.iter_rows(min_row=2, values_only=True):
        text = lat(row[0] if row else '')
        if not text:
            continue
        upper = text.upper()
        if upper in {'DA', 'NE'} or upper.startswith('NP') or upper.startswith('NAPOMENA'):
            continue
        if upper.startswith('NAZIV SUBJEKTA') or upper == 'PITANJE':
            continue
        if upper.startswith('ČEKIRATI') or upper.startswith('CHEKIRATI'):
            continue
        if upper.startswith('PRIME'):
            continue
        # section headers vs questions
        if text.endswith('?') or len(text) > 40 or text[:1].islower():
            n += 1
            lines.append(f'{n}. {text}')
        else:
            lines.append('')
            lines.append(text.upper())
            lines.append('')
    return '\n'.join(lines).strip() + '\n'


def build_entries():
    return [
        {
            'code': 'IN-01',
            'kind': 'IN',
            'title': 'Tipične neusaglašenosti i kaznene mere',
            'source_filename': 'Report Sloy B.xlsx',
            'body': fines_body(),
            'is_fixed': True,
            'note': '',
            'sort_order': 10,
        },
        {
            'code': 'IN-02',
            'kind': 'IN',
            'title': 'Kontrolna lista — sledljivost',
            'source_filename': 'Kontrolne liste digital empty.xlsx',
            'body': checklist_body('Следљивост'),
            'is_fixed': True,
            'note': '',
            'sort_order': 20,
        },
        {
            'code': 'IN-03',
            'kind': 'IN',
            'title': 'Kontrolna lista — higijena',
            'source_filename': 'Kontrolne liste digital empty.xlsx',
            'body': checklist_body('Хигијена'),
            'is_fixed': True,
            'note': '',
            'sort_order': 30,
        },
        {
            'code': 'IN-04',
            'kind': 'IN',
            'title': 'Kontrolna lista — HACCP',
            'source_filename': 'Kontrolne liste digital empty.xlsx',
            'body': checklist_body('HACCP'),
            'is_fixed': True,
            'note': '',
            'sort_order': 40,
        },
    ]


if __name__ == '__main__':
    for e in build_entries():
        print(e['code'], e['title'], 'chars', len(e['body']))
